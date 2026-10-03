"""Tests for the V2 report Markdown renderer and PPTX renderer."""

from __future__ import annotations

from pathlib import Path

import pytest

from app.domain.generation_documents import Presentation, Report, Slide
from app.infrastructure.pptx_renderer import (
    PptxRendererError,
    render_presentation_pptx,
    slug_filename,
)
from app.templates.report import render_report_markdown


def _report() -> Report:
    return Report(
        title="Quarterly Review",
        summary="Summary paragraph.",
        sections=[
            {
                "heading": "Background",
                "paragraphs": ["First paragraph.", "Second paragraph."],
                "bullets": ["point one", "point two"],
                "references": ["source-a"],
            },
            {
                "heading": "Numbers",
                "paragraphs": [],
                "bullets": [],
                "table": {"headers": ["a", "b"], "rows": [["1", "2"]]},
                "references": [],
            },
        ],
    )


def test_report_markdown_deterministic() -> None:
    first = render_report_markdown(_report())
    second = render_report_markdown(_report())

    assert first == second
    assert first.startswith("# Quarterly Review\n\nSummary paragraph.")
    assert "## Background" in first
    assert "First paragraph." in first
    assert "- point one" in first
    assert "| a | b |" in first
    assert "| 1 | 2 |" in first
    assert "- [1] source-a" in first
    assert "### References" in first


def test_report_markdown_skips_empty_parts() -> None:
    report = Report(
        title="T",
        summary="S.",
        sections=[{"heading": "Only", "paragraphs": ["Body."]}],
    )

    rendered = render_report_markdown(report)

    assert "### References" not in rendered
    assert "## Only" in rendered


def _presentation() -> Presentation:
    return Presentation(
        title="Deck",
        slides=[
            {"title": "Intro", "bullets": ["one", "two"], "speaker_notes": "say hi"},
            {"title": "Outro", "bullets": ["three"], "speaker_notes": None},
        ],
    )


def _destination(tmp_path: Path, name: str = "deck.pptx") -> tuple[Path, Path]:
    root = tmp_path / "artifacts"
    return root, root / name


def test_pptx_generated_file_exists_and_valid(tmp_path: Path) -> None:
    root, destination = _destination(tmp_path)

    returned = render_presentation_pptx(_presentation(), destination, artifact_root=root)

    assert returned == destination.resolve()
    assert destination.exists()


def test_pptx_slide_count_titles_and_bullets(tmp_path: Path) -> None:
    root, destination = _destination(tmp_path)
    render_presentation_pptx(_presentation(), destination, artifact_root=root)

    from pptx import Presentation as PptxReopen

    deck = PptxReopen(str(destination))
    titles = [
        slide.shapes.title.text if slide.shapes.title is not None else ""
        for slide in deck.slides
    ]

    assert titles == ["Deck", "Intro", "Outro"]
    content = deck.slides[1].placeholders[1].text
    assert "one" in content
    assert "two" in content


def test_pptx_speaker_notes_enabled(tmp_path: Path) -> None:
    root, destination = _destination(tmp_path)
    render_presentation_pptx(_presentation(), destination, artifact_root=root)

    from pptx import Presentation as PptxReopen

    deck = PptxReopen(str(destination))

    assert "say hi" in deck.slides[1].notes_slide.notes_text_frame.text


def test_pptx_speaker_notes_omitted_when_disabled(tmp_path: Path) -> None:
    root, destination = _destination(tmp_path)
    render_presentation_pptx(
        _presentation(), destination, artifact_root=root, include_notes=False
    )

    from pptx import Presentation as PptxReopen

    deck = PptxReopen(str(destination))

    assert "say hi" not in deck.slides[1].notes_slide.notes_text_frame.text


def test_pptx_rejects_destination_outside_artifact_root(tmp_path: Path) -> None:
    root = tmp_path / "artifacts"
    outside = tmp_path / "elsewhere" / "deck.pptx"

    with pytest.raises(PptxRendererError):
        render_presentation_pptx(_presentation(), outside, artifact_root=root)


def test_slug_filename() -> None:
    assert slug_filename("Quarterly Review: Q1/Q2?") == "Quarterly-Review-Q1Q2"
    assert slug_filename("   ") == "presentation"


def test_pptx_deterministic_structure(tmp_path: Path) -> None:
    from pptx import Presentation as PptxReopen

    root, first_path = _destination(tmp_path, "first.pptx")
    _, second_path = _destination(tmp_path, "second.pptx")
    render_presentation_pptx(_presentation(), first_path, artifact_root=root)
    render_presentation_pptx(_presentation(), second_path, artifact_root=root)

    def titles(path: Path) -> list[str]:
        deck = PptxReopen(str(path))
        return [
            slide.shapes.title.text if slide.shapes.title is not None else ""
            for slide in deck.slides
        ]

    assert titles(first_path) == titles(second_path) == ["Deck", "Intro", "Outro"]


def test_slide_model_allows_empty_bullets() -> None:
    slide = Slide(title="Divider", bullets=[])

    assert slide.bullets == []
